"""Unit tests for document text extraction.

Documents are built in-memory rather than committed as fixtures, so the tests
exercise the real pypdf / python-docx code paths without binary test data.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from src.domains.resume.parsing import (
    CONTENT_TYPE_DOCX,
    CONTENT_TYPE_PDF,
    LegacyDocNotSupported,
    NoTextContent,
    UnparseableDocument,
    UnsupportedFileType,
    detect_content_type,
    extract_text,
    resolve_content_type,
)

LONG_TEXT = (
    "Ada Lovelace — Final-year Computer Science student at IIT Bombay. "
    "Backend engineering intern at Acme Corp working on Python and PostgreSQL. "
    "Built a toy compiler in Rust and a distributed cache in Go."
)


def _docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for cell_index, value in enumerate(row):
                table.rows[row_index].cells[cell_index].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _empty_pdf_bytes(pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_extracts_docx_paragraphs() -> None:
    result = extract_text(_docx_bytes([LONG_TEXT]), CONTENT_TYPE_DOCX)

    assert "Ada Lovelace" in result
    assert "IIT Bombay" in result


def test_extracts_docx_table_cells() -> None:
    """Resumes lay dates and roles out in tables; skipping them drops whole sections."""
    result = extract_text(
        _docx_bytes(
            [LONG_TEXT, "Experience"],
            table_rows=[["Acme Corp", "Backend Intern", "Jun 2025"]],
        ),
        CONTENT_TYPE_DOCX,
    )

    assert "Acme Corp" in result
    assert "Backend Intern" in result
    assert "Jun 2025" in result


def test_blank_pdf_raises_no_text_content() -> None:
    """A PDF with no text layer is a scan — a distinct, actionable failure."""
    with pytest.raises(NoTextContent):
        extract_text(_empty_pdf_bytes(), CONTENT_TYPE_PDF)


def test_corrupt_pdf_raises_unparseable() -> None:
    with pytest.raises(UnparseableDocument):
        extract_text(b"this is definitely not a pdf", CONTENT_TYPE_PDF)


def test_corrupt_docx_raises_unparseable() -> None:
    with pytest.raises(UnparseableDocument):
        extract_text(b"not a docx either", CONTENT_TYPE_DOCX)


def test_unsupported_content_type_is_rejected() -> None:
    with pytest.raises(UnsupportedFileType):
        extract_text(b"anything", "text/plain")


def test_near_empty_document_raises_no_text_content() -> None:
    with pytest.raises(NoTextContent):
        extract_text(_docx_bytes(["Hi"]), CONTENT_TYPE_DOCX)


@pytest.mark.parametrize(
    ("filename", "declared", "expected"),
    [
        ("resume.pdf", None, CONTENT_TYPE_PDF),
        ("Resume.PDF", "application/octet-stream", CONTENT_TYPE_PDF),
        ("cv.docx", None, CONTENT_TYPE_DOCX),
        ("no-extension", CONTENT_TYPE_PDF, CONTENT_TYPE_PDF),
    ],
)
def test_content_type_resolution(filename: str, declared: str | None, expected: str) -> None:
    """The extension wins over the browser-declared type, which is client-controlled."""
    assert resolve_content_type(filename, declared) == expected


def test_unsupported_extension_is_rejected() -> None:
    with pytest.raises(UnsupportedFileType):
        resolve_content_type("resume.txt", "text/plain")


# --------------------------------------------------------------------------
# Content-based detection — the upload path's actual authority
# --------------------------------------------------------------------------


def test_detect_content_type_reads_the_pdf_magic_number() -> None:
    assert detect_content_type(io.BytesIO(b"%PDF-1.7\nrest of file")) == CONTENT_TYPE_PDF


def test_detect_content_type_accepts_a_real_docx() -> None:
    assert detect_content_type(io.BytesIO(_docx_bytes(["Ada Lovelace"]))) == CONTENT_TYPE_DOCX


def test_detect_content_type_rejects_a_renamed_text_file() -> None:
    """The name and the browser's Content-Type both say PDF; the bytes do not.
    This is the case extension-based resolution cannot see."""
    with pytest.raises(UnsupportedFileType):
        detect_content_type(io.BytesIO(b"Ada Lovelace, backend engineer, not a PDF at all"))


def test_detect_content_type_rejects_legacy_doc_with_its_own_error() -> None:
    """OLE2 gets a distinct error so the student is told to re-save rather than
    being told .doc is simply unsupported."""
    with pytest.raises(LegacyDocNotSupported) as excinfo:
        detect_content_type(io.BytesIO(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64))
    assert "Legacy .doc" in excinfo.value.message


def test_legacy_doc_error_is_an_unsupported_file_type() -> None:
    """Callers that only catch the general error must still reject it."""
    assert issubclass(LegacyDocNotSupported, UnsupportedFileType)


def test_detect_content_type_rejects_a_zip_that_is_not_a_word_document() -> None:
    """Every OOXML format shares the zip magic number, so the magic alone cannot
    tell a .docx from an .xlsx."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
    buffer.seek(0)

    with pytest.raises(UnsupportedFileType):
        detect_content_type(buffer)


def test_detect_content_type_rewinds_the_stream() -> None:
    """The caller streams the same object straight to object storage, so a
    consumed or half-read stream would upload a truncated file."""
    data = _docx_bytes(["Ada Lovelace"])
    fileobj = io.BytesIO(data)

    detect_content_type(fileobj)

    assert fileobj.tell() == 0
    assert fileobj.read() == data
