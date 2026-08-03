"""Text extraction from PDF and DOCX bytes.

Every failure mode here is explicit and typed, because "the file could not be
read" and "the file was read but contained no text" need different messages to
the student — the first is a broken or protected file, the second is almost
always a scanned image, which no amount of retrying will fix.

All errors raised here are deterministic: the worker fails the job immediately
rather than retrying a file that will parse identically every time.
"""

from __future__ import annotations

import io
import zipfile
from typing import BinaryIO

import structlog
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from src.core.exceptions import AppError

logger = structlog.get_logger(__name__)

CONTENT_TYPE_PDF = "application/pdf"
CONTENT_TYPE_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SUPPORTED_CONTENT_TYPES = {CONTENT_TYPE_PDF, CONTENT_TYPE_DOCX}
SUPPORTED_EXTENSIONS = {".pdf": CONTENT_TYPE_PDF, ".docx": CONTENT_TYPE_DOCX}

# Below this, the "document" is almost certainly a scanned image with no text
# layer. Sending it to the LLM would bill a request to extract nothing.
MIN_MEANINGFUL_CHARS = 100


class UnsupportedFileType(AppError):
    """The uploaded file is not a PDF or DOCX."""

    status_code = 415
    code = "UNSUPPORTED_FILE_TYPE"


class LegacyDocNotSupported(UnsupportedFileType):
    """An OLE2 `.doc`. Detected separately so the student gets a fix, not a shrug.

    Worth its own error because the remedy is specific and trivial — re-save as
    PDF or DOCX — and because `python-docx` reads OOXML only, so accepting it
    would mean a 30-second wait ending in a generic worker failure.
    """

    code = "LEGACY_DOC_NOT_SUPPORTED"


class UnparseableDocument(AppError):
    """The file could not be read — corrupt, encrypted, or not really its declared type."""

    status_code = 422
    code = "UNPARSEABLE_DOCUMENT"


class NoTextContent(AppError):
    """The file parsed but held no extractable text (typically a scan)."""

    status_code = 422
    code = "NO_TEXT_CONTENT"


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        # Encrypted PDFs parse structurally but yield empty pages, which would
        # otherwise surface as the misleading "no text" error.
        if reader.is_encrypted:
            raise UnparseableDocument(
                "This PDF is password protected. Upload an unprotected copy."
            )
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except UnparseableDocument:
        raise
    except (PdfReadError, ValueError, OSError) as exc:
        logger.warning("pdf_parse_failed", error=str(exc))
        raise UnparseableDocument("This PDF could not be read. It may be corrupt.") from exc


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    # A DOCX is a zip archive: a file that isn't one surfaces as BadZipFile,
    # which is not an OSError subclass and would otherwise escape untyped.
    except (PackageNotFoundError, zipfile.BadZipFile, ValueError, OSError) as exc:
        logger.warning("docx_parse_failed", error=str(exc))
        raise UnparseableDocument("This DOCX could not be read. It may be corrupt.") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    # Resumes frequently lay out dates and roles in tables, whose text is not
    # in `paragraphs` — skipping them silently drops whole experience sections.
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def extract_text(data: bytes, content_type: str) -> str:
    """Extract plain text, raising a typed error for every failure mode."""
    if content_type == CONTENT_TYPE_PDF:
        text = _extract_pdf(data)
    elif content_type == CONTENT_TYPE_DOCX:
        text = _extract_docx(data)
    else:
        raise UnsupportedFileType(f"Unsupported content type '{content_type}'")

    cleaned = text.strip()
    if len(cleaned) < MIN_MEANINGFUL_CHARS:
        raise NoTextContent(
            "No readable text was found. If this is a scanned resume, upload a "
            "text-based PDF or DOCX instead."
        )
    return cleaned


def resolve_content_type(filename: str, declared_content_type: str | None) -> str:
    """Decide the content type from the filename extension.

    A *hint* only, kept for callers that have a name but no bytes. It is not
    the upload path's authority — `detect_content_type` is, because both the
    extension and the browser-declared type are client-controlled and a
    `.pdf` that is really a zip must not reach the worker as a PDF.
    """
    lowered = filename.lower()
    for extension, content_type in SUPPORTED_EXTENSIONS.items():
        if lowered.endswith(extension):
            return content_type

    if declared_content_type in SUPPORTED_CONTENT_TYPES:
        return declared_content_type

    raise UnsupportedFileType("Only PDF and DOCX resumes are supported")


# --------------------------------------------------------------------------
# Content-based type detection
# --------------------------------------------------------------------------

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"
# OLE2 compound-file header — the container legacy `.doc`, `.xls` and `.ppt`
# all share.
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

_SNIFF_BYTES = 8

LEGACY_DOC_MESSAGE = "Legacy .doc isn't supported — save as PDF or DOCX"


def _is_wordprocessing_package(fileobj: BinaryIO) -> bool:
    """Is this zip specifically a Word document, not just any zip?

    OOXML formats are all zips, so the magic number alone cannot tell a `.docx`
    from a `.xlsx` or an ordinary archive. `word/document.xml` is the part that
    makes a package a wordprocessing one. Reading the central directory is
    cheap — `zipfile` seeks to the end rather than scanning the whole file.
    """
    try:
        with zipfile.ZipFile(fileobj) as archive:
            return "word/document.xml" in archive.namelist()
    except (zipfile.BadZipFile, OSError, ValueError):
        return False
    finally:
        fileobj.seek(0)


def detect_content_type(fileobj: BinaryIO) -> str:
    """Determine the content type from the bytes themselves.

    This is the upload path's authority. The filename extension and the
    browser's `Content-Type` are both attacker-controlled, so neither is
    consulted: a file is a PDF because it starts with `%PDF-`, and a DOCX
    because it is a zip containing `word/document.xml`.

    Leaves `fileobj` rewound to 0 so the caller can stream it straight to
    object storage.
    """
    head = fileobj.read(_SNIFF_BYTES)
    fileobj.seek(0)

    if head.startswith(_PDF_MAGIC):
        return CONTENT_TYPE_PDF

    if head.startswith(_OLE2_MAGIC):
        raise LegacyDocNotSupported(LEGACY_DOC_MESSAGE)

    if head.startswith(_ZIP_MAGIC):
        if _is_wordprocessing_package(fileobj):
            return CONTENT_TYPE_DOCX
        raise UnsupportedFileType(
            "That archive is not a Word document. Only PDF and DOCX resumes are supported."
        )

    raise UnsupportedFileType("Only PDF and DOCX resumes are supported")
